from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

COLORS = {
    "navy": RGBColor(13, 45, 97),
    "blue": RGBColor(24, 95, 165),
    "light_blue": RGBColor(230, 242, 255),
    "mid_blue": RGBColor(92, 149, 214),
    "dark": RGBColor(30, 40, 58),
    "gray": RGBColor(96, 108, 128),
    "light_gray": RGBColor(245, 248, 252),
    "white": RGBColor(255, 255, 255),
}


def set_slide_bg(slide, rgb):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def add_top_bar(slide, label="Faculty Job Talk"):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, Inches(0.45))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS["navy"]
    bar.line.fill.background()
    text = bar.text_frame
    text.clear()
    p = text.paragraphs[0]
    p.text = label
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLORS["white"]
    p.alignment = PP_ALIGN.LEFT


def add_title_block(slide, title, subtitle=""):
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.8), Inches(0.9))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Calibri"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = COLORS["navy"]

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.82), Inches(1.55), Inches(11.0), Inches(0.6))
        stf = sub_box.text_frame
        stf.clear()
        sp = stf.paragraphs[0]
        sp.text = subtitle
        sp.font.name = "Calibri"
        sp.font.size = Pt(21)
        sp.font.color.rgb = COLORS["blue"]


def add_bullets(slide, x, y, w, h, items, font_size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Calibri"
        p.font.size = Pt(font_size)
        p.font.color.rgb = COLORS["dark"]
        p.level = 0
        p.space_after = Pt(6)
        p.space_before = Pt(2)


def add_card(slide, x, y, w, h, title, body, icon_shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    card = slide.shapes.add_shape(icon_shape, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = COLORS["white"]
    card.line.color.rgb = COLORS["mid_blue"]
    card.line.width = Pt(1.5)

    tf = card.text_frame
    tf.clear()
    p1 = tf.paragraphs[0]
    p1.text = title
    p1.font.name = "Calibri"
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = COLORS["navy"]
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.name = "Calibri"
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLORS["gray"]


def add_timeline(slide, labels):
    line_y = Inches(3.6)
    slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(1.2), line_y, Inches(12.2), line_y
    ).line.color.rgb = COLORS["mid_blue"]

    x_positions = [1.2 + i * (11.0 / (len(labels) - 1)) for i in range(len(labels))]
    for i, (x, label) in enumerate(zip(x_positions, labels)):
        node = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - 0.17), Inches(3.42), Inches(0.34), Inches(0.34))
        node.fill.solid()
        node.fill.fore_color.rgb = COLORS["blue"] if i < len(labels) - 1 else COLORS["navy"]
        node.line.fill.background()

        box = slide.shapes.add_textbox(Inches(x - 0.9), Inches(2.2 if i % 2 == 0 else 3.95), Inches(1.8), Inches(0.9))
        tf = box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = label
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Calibri"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLORS["dark"]


def add_section_title(slide, title):
    add_top_bar(slide)
    add_title_block(slide, title)


def add_metric_card(slide, x, y, title, value="XX"):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.4), Inches(1.35))
    card.fill.solid()
    card.fill.fore_color.rgb = COLORS["light_blue"]
    card.line.color.rgb = COLORS["mid_blue"]
    tf = card.text_frame
    tf.clear()
    p1 = tf.paragraphs[0]
    p1.text = value
    p1.font.name = "Calibri"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = COLORS["navy"]
    p1.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.name = "Calibri"
    p2.font.size = Pt(12)
    p2.font.color.rgb = COLORS["dark"]
    p2.alignment = PP_ALIGN.CENTER


def add_research_slide(slide, project_title):
    add_top_bar(slide)
    add_title_block(slide, project_title, "Selected Research Contribution")

    left_x, right_x = 0.7, 6.9
    add_card(slide, left_x, 2.0, 5.7, 1.1, "Problem", "Define challenge, gap, and domain relevance.")
    add_card(slide, left_x, 3.2, 5.7, 1.15, "Methodology", "Pipeline, datasets, and learning strategy.")
    add_card(slide, left_x, 4.45, 5.7, 1.15, "Model Architecture", "Backbone, modules, and training objective.")

    add_card(slide, right_x, 2.0, 5.7, 1.4, "Experimental Results", "Accuracy / F1 / AUC placeholders and baseline gains.")
    add_card(slide, right_x, 3.55, 5.7, 1.0, "Key Contributions", "Novelty, interpretability, and clinical/real impact.")
    add_card(slide, right_x, 4.7, 5.7, 1.0, "Future Applications", "Deployment pathway, transfer learning, and extension.")

    # Architecture diagram area
    arch = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(6.0), Inches(4.3), Inches(1.2))
    arch.fill.solid()
    arch.fill.fore_color.rgb = COLORS["light_gray"]
    arch.line.color.rgb = COLORS["mid_blue"]
    tf = arch.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "AI Architecture Diagram Placeholder"
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.italic = True
    p.font.color.rgb = COLORS["gray"]
    p.alignment = PP_ALIGN.CENTER


def build_presentation(output_path):
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank = prs.slide_layouts[6]

    # 1 Title
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_top_bar(slide, "Department of Computer Science and Information Engineering (CSIE) | Chang Gung University")

    add_title_block(slide, "Faculty Job Talk", "Dr. Ammar Amjad")
    info = slide.shapes.add_textbox(Inches(0.85), Inches(2.1), Inches(7.4), Inches(1.5))
    tf = info.text_frame
    tf.clear()
    p1 = tf.paragraphs[0]
    p1.text = "Postdoctoral Researcher"
    p1.font.name = "Calibri"
    p1.font.size = Pt(24)
    p1.font.color.rgb = COLORS["blue"]
    p2 = tf.add_paragraph()
    p2.text = "National Yang Ming Chiao Tung University (NYCU)"
    p2.font.name = "Calibri"
    p2.font.size = Pt(20)
    p2.font.color.rgb = COLORS["dark"]

    # Abstract AI background panel
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.0), Inches(1.2), Inches(4.9), Inches(5.8))
    panel.fill.solid()
    panel.fill.fore_color.rgb = COLORS["light_blue"]
    panel.line.color.rgb = COLORS["mid_blue"]
    nodes = [(8.5, 2.0), (10.2, 1.9), (11.7, 2.5), (9.2, 3.2), (10.8, 3.7), (12.0, 4.4), (9.6, 5.1), (11.1, 5.7)]
    for i, (x, y) in enumerate(nodes):
        n = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.32), Inches(0.32))
        n.fill.solid()
        n.fill.fore_color.rgb = COLORS["blue"] if i % 2 == 0 else COLORS["mid_blue"]
        n.line.fill.background()
    for i in range(len(nodes) - 1):
        conn = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(nodes[i][0] + 0.15),
            Inches(nodes[i][1] + 0.15),
            Inches(nodes[i + 1][0] + 0.15),
            Inches(nodes[i + 1][1] + 0.15),
        )
        conn.line.color.rgb = COLORS["mid_blue"]
        conn.line.width = Pt(1.2)

    # 2 About Me
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_section_title(slide, "About Me")
    about_items = [
        "Academic background in AI, Machine Learning, and Computer Vision",
        "Education: BSc -> MSc -> PhD in Computer Science related disciplines",
        "Professional experience across academia and applied AI projects",
        "Current position: Postdoctoral Researcher at NYCU",
        "Research interests: NLP, Medical AI, Explainable & Generative AI",
        "Career highlights: publications, supervision, and collaborations",
    ]
    add_bullets(slide, 0.9, 2.0, 7.6, 4.8, about_items, font_size=17)
    add_card(slide, 8.8, 2.0, 3.8, 1.2, "Research Focus", "Human-centered, trustworthy, and translational AI.")
    add_card(slide, 8.8, 3.45, 3.8, 1.2, "Teaching Vision", "Engaging classrooms with project-driven learning.")
    add_card(slide, 8.8, 4.9, 3.8, 1.2, "Faculty Goal", "Build impactful AI programs and mentor talent.")

    # 3 Academic Journey
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_section_title(slide, "Academic Journey")
    add_timeline(
        slide,
        [
            "Bachelor's\nDegree",
            "Master's\nDegree",
            "PhD",
            "Research\nExperience",
            "Postdoctoral\nResearch",
            "Faculty\nCareer Goal",
        ],
    )

    # 4 Teaching Philosophy
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_section_title(slide, "Teaching Philosophy")
    philosophy = [
        ("Student-Centered Learning", "Personalized guidance and inclusive learning outcomes."),
        ("Active Learning", "In-class problem solving, peer discussion, and feedback loops."),
        ("Project-Based Learning", "Real-world AI/CS projects with measurable milestones."),
        ("Research-Integrated Teaching", "Latest papers and lab experience in coursework."),
        ("AI-Assisted Education", "Use of coding assistants and adaptive learning analytics."),
        ("Industry Collaboration", "Guest lectures and co-designed capstone experiences."),
    ]
    x0, y0 = 0.8, 2.0
    for i, (t, b) in enumerate(philosophy):
        add_card(slide, x0 + (i % 3) * 4.2, y0 + (i // 3) * 2.2, 3.9, 1.9, t, b)

    # 5 Courses I Can Teach
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_section_title(slide, "Courses I Can Teach")
    ug_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(5.9), Inches(4.8))
    ug_box.fill.solid()
    ug_box.fill.fore_color.rgb = COLORS["light_blue"]
    ug_box.line.color.rgb = COLORS["mid_blue"]
    g_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(2.0), Inches(5.6), Inches(4.8))
    g_box.fill.solid()
    g_box.fill.fore_color.rgb = COLORS["light_gray"]
    g_box.line.color.rgb = COLORS["mid_blue"]
    add_bullets(
        slide,
        1.1,
        2.25,
        5.2,
        4.4,
        [
            "Undergraduate",
            "Programming, Python Programming",
            "Data Structures, Algorithms",
            "Artificial Intelligence, Machine Learning",
            "Computer Vision, Image Processing",
            "Database Systems, Operating Systems",
        ],
        font_size=16,
    )
    add_bullets(
        slide,
        7.2,
        2.25,
        5.0,
        4.4,
        [
            "Graduate",
            "Deep Learning, Advanced Machine Learning",
            "Natural Language Processing, Large Language Models",
            "Computer Vision, Medical AI",
            "Explainable AI, Generative AI",
        ],
        font_size=16,
    )

    # 6 Teaching Demonstration
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_section_title(slide, "Teaching Demonstration: Introduction to Digital Image Processing")
    topics = [
        "What is an image?",
        "Pixels",
        "RGB Color Model",
        "Grayscale images",
        "Image filtering",
        "Edge detection",
        "CNN applications",
        "Medical imaging examples",
    ]
    add_bullets(slide, 0.9, 2.0, 5.2, 4.9, topics, font_size=15)
    # Simple educational diagram
    img_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.3), Inches(2.0), Inches(6.3), Inches(4.8))
    img_box.fill.solid()
    img_box.fill.fore_color.rgb = COLORS["light_gray"]
    img_box.line.color.rgb = COLORS["mid_blue"]
    tf = img_box.text_frame
    tf.text = "Image Processing Pipeline Diagram\nInput -> Filter -> Edge Map -> CNN -> Clinical Insight"
    tf.paragraphs[0].font.name = "Calibri"
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.color.rgb = COLORS["navy"]
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    # 7 Teaching Experience
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_section_title(slide, "Teaching Experience")
    blocks = [
        ("Courses Taught", "AI, Machine Learning, Programming, and Image Analysis related courses."),
        ("Student Supervision", "Supervised undergraduate projects and graduate theses."),
        ("Practical Projects", "Hands-on team projects with real datasets."),
        ("Programming Assignments", "Scaffolded assignments from basics to model deployment."),
        ("Interactive Methods", "Flipped classroom, peer review, and coding clinics."),
    ]
    for i, (t, b) in enumerate(blocks):
        add_card(slide, 0.9 + (i % 2) * 6.3, 2.0 + (i // 2) * 1.6, 5.9, 1.4, t, b)

    # 8 Research Overview
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_section_title(slide, "Research Overview")
    areas = [
        "Natural Language Processing",
        "Computer Vision",
        "Deep Learning",
        "Speech Processing",
        "Medical AI",
        "Explainable AI",
        "Generative AI",
        "Multimodal AI",
    ]
    for i, area in enumerate(areas):
        x = 0.9 + (i % 4) * 3.1
        y = 2.2 + (i // 4) * 2.3
        bubble = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(2.7), Inches(1.8))
        bubble.fill.solid()
        bubble.fill.fore_color.rgb = COLORS["light_blue"] if i % 2 == 0 else COLORS["light_gray"]
        bubble.line.color.rgb = COLORS["mid_blue"]
        tf = bubble.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = area
        p.font.name = "Calibri"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLORS["dark"]
        p.alignment = PP_ALIGN.CENTER

    # 9 Research Impact
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_section_title(slide, "Research Impact")
    metrics = [
        "Total Publications",
        "SCI Journal Papers",
        "Conference Papers",
        "Citations",
        "h-index",
        "i10-index",
    ]
    for i, m in enumerate(metrics):
        add_metric_card(slide, 0.9 + (i % 3) * 4.1, 2.3 + (i // 3) * 1.9, m, "XX")

    chart_data = ChartData()
    chart_data.categories = ["Y1", "Y2", "Y3", "Y4", "Y5"]
    chart_data.add_series("Publications", (2, 4, 6, 8, 10))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS, Inches(8.2), Inches(5.6), Inches(4.5), Inches(1.6), chart_data
    ).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM

    # 10-14 Selected Research Contributions
    titles = [
        "Selected Research Contribution 1: NLP for Clinical Text Analytics",
        "Selected Research Contribution 2: Vision Transformers for Medical Imaging",
        "Selected Research Contribution 3: Explainable AI for Diagnostic Support",
        "Selected Research Contribution 4: Multimodal Speech-Language Understanding",
        "Selected Research Contribution 5: Generative AI for Data Augmentation",
    ]
    for t in titles:
        slide = prs.slides.add_slide(blank)
        set_slide_bg(slide, COLORS["white"])
        add_research_slide(slide, t)

    # 15 Current Research Projects
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_section_title(slide, "Current Research Projects")
    current_projects = [
        "AI for Healthcare",
        "Speech Emotion Recognition",
        "Natural Language Processing",
        "Medical Image Analysis",
        "Industry AI Projects",
        "High Performance Computing",
    ]
    for i, item in enumerate(current_projects):
        add_card(slide, 0.9 + (i % 3) * 4.1, 2.1 + (i // 3) * 2.3, 3.8, 1.9, item, "Project summary placeholder.")

    # 16 Future Research Directions
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_section_title(slide, "Future Research Directions")
    roadmap = [
        "Multimodal AI",
        "Large Language Models",
        "Explainable AI",
        "Trustworthy AI",
        "Medical AI",
        "Edge AI",
        "AI for Healthcare",
        "Human-Centered AI",
        "Industry Collaboration",
        "International Collaboration",
    ]
    for i, step in enumerate(roadmap):
        x = 0.8 + i * 1.2
        y = 3.0 + (i % 2) * 1.2
        node = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(1.15), Inches(0.8))
        node.fill.solid()
        node.fill.fore_color.rgb = COLORS["mid_blue"] if i % 2 == 0 else COLORS["blue"]
        node.line.fill.background()
        tf = node.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = step
        p.font.name = "Calibri"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLORS["white"]
        p.alignment = PP_ALIGN.CENTER

    # 17 Academic Service
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_section_title(slide, "Academic Service")
    services = [
        "Journal Reviewer",
        "Conference Reviewer",
        "Editorial Activities",
        "Research Collaboration",
        "Student Mentoring",
        "Professional Memberships",
    ]
    for i, srv in enumerate(services):
        add_card(slide, 1.0 + (i % 3) * 4.1, 2.2 + (i // 3) * 2.1, 3.7, 1.7, srv, "Service portfolio placeholder.")

    # 18 Why Chang Gung University
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_section_title(slide, "Why Chang Gung University")
    why = [
        "Alignment with CSIE research strengths",
        "Potential collaborations across AI and engineering",
        "Medical AI opportunities with clinical ecosystem",
        "Interdisciplinary research and translational impact",
        "Strong teaching contributions and curriculum innovation",
        "Student mentorship with long-term research vision",
    ]
    add_bullets(slide, 0.9, 2.0, 7.4, 4.8, why, font_size=16)
    add_card(slide, 8.5, 2.1, 4.0, 4.6, "Strategic Fit", "Collaborative ecosystem to build a globally competitive AI program.")

    # 19 Five-Year Vision
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_section_title(slide, "Five-Year Vision")
    vision = [
        "Research Lab Development",
        "Graduate Student Supervision",
        "International Collaboration",
        "Research Funding",
        "Industry Partnerships",
        "High-impact Publications",
    ]
    for i, goal in enumerate(vision):
        add_card(slide, 0.9 + (i % 3) * 4.1, 2.2 + (i // 3) * 2.2, 3.8, 1.8, goal, "Milestones and KPIs placeholder.")

    # 20 Thank You
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, COLORS["white"])
    add_top_bar(slide, "Thank You")
    add_title_block(slide, "Thank You", "Questions & Discussion")
    closing = slide.shapes.add_textbox(Inches(3.0), Inches(3.0), Inches(7.3), Inches(2.0))
    tf = closing.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "Dr. Ammar Amjad\nPostdoctoral Researcher, NYCU"
    p.font.name = "Calibri"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLORS["navy"]
    p.alignment = PP_ALIGN.CENTER

    prs.save(output_path)


if __name__ == "__main__":
    build_presentation("/workspace/faculty_job_talk_cgu_csie_ammar_amjad.pptx")
    print("Created /workspace/faculty_job_talk_cgu_csie_ammar_amjad.pptx")
